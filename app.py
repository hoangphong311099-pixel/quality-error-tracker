import csv
import os
import shutil
import sqlite3
import tkinter as tk
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox

import customtkinter as ctk
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from matplotlib.figure import Figure
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from PIL import Image as PILImage, ImageTk

ROOT_DIR = Path(__file__).resolve().parent
DATA_DIR = ROOT_DIR / "data"
IMAGE_DIR = DATA_DIR / "images"
BACKUP_DIR = ROOT_DIR / "backups"
EXPORT_DIR = ROOT_DIR / "exports"
DB_PATH = DATA_DIR / "error_tracker.db"

os.makedirs(DATA_DIR, exist_ok=True)
os.makedirs(IMAGE_DIR, exist_ok=True)
os.makedirs(BACKUP_DIR, exist_ok=True)
os.makedirs(EXPORT_DIR, exist_ok=True)


def status_color(status: str):
    color_map = {
        "Open": "#2E7DFF",
        "Processed": "#7C4DFF",
        "Processing": "#FFB300",
        "Closed": "#2E7D32",
    }
    return color_map.get(status, "#607D8B")


class Database:
    def __init__(self, db_path: str = str(DB_PATH)):
        self.db_path = db_path
        self.init_db()

    def connect(self):
        conn = sqlite3.connect(self.db_path)
        conn.row_factory = sqlite3.Row
        return conn

    def init_db(self):
        conn = self.connect()

        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT UNIQUE NOT NULL,
                password TEXT NOT NULL,
                role TEXT NOT NULL CHECK(role IN ('Admin', 'Engineer', 'Viewer')),
                full_name TEXT,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP
            );

            CREATE TABLE IF NOT EXISTS defects (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                defect_code TEXT,
                model TEXT,
                line_name TEXT,
                defect_name TEXT,
                responsible TEXT,
                status TEXT,
                defect_date TEXT,
                description TEXT,
                root_cause TEXT,
                improvement_method TEXT,
                before_image_path TEXT,
                after_image_path TEXT,
                created_by INTEGER,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                updated_at TEXT DEFAULT CURRENT_TIMESTAMP
            );
            """
        )

        cols = conn.execute("PRAGMA table_info(defects)").fetchall()
        existing = [c[1] for c in cols]

        if "improvement_method" not in existing:
            conn.execute("ALTER TABLE defects ADD COLUMN improvement_method TEXT")
        if "before_image_path" not in existing:
            conn.execute("ALTER TABLE defects ADD COLUMN before_image_path TEXT")
        if "after_image_path" not in existing:
            conn.execute("ALTER TABLE defects ADD COLUMN after_image_path TEXT")

        default_users = [
            ("admin", "admin123", "Admin", "Administrator"),
            ("engineer", "engineer123", "Engineer", "Process Engineer"),
            ("viewer", "viewer123", "Viewer", "Viewer User"),
        ]

        for username, password, role, full_name in default_users:
            conn.execute(
                "INSERT OR IGNORE INTO users (username, password, role, full_name) VALUES (?, ?, ?, ?)",
                (username, password, role, full_name),
            )

        conn.commit()
        conn.close()

    def authenticate(self, username, password):
        conn = self.connect()
        user = conn.execute(
            "SELECT * FROM users WHERE username = ? AND password = ?",
            (username, password),
        ).fetchone()
        conn.close()
        return dict(user) if user else None

    def get_user_by_id(self, user_id):
        conn = self.connect()
        row = conn.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
        conn.close()
        return dict(row) if row else None

    def get_all_users(self):
        conn = self.connect()
        rows = conn.execute("SELECT * FROM users ORDER BY id").fetchall()
        conn.close()
        return [dict(r) for r in rows]

    def insert_user(self, username, password, role, full_name):
        conn = self.connect()
        try:
            conn.execute(
                "INSERT INTO users (username, password, role, full_name) VALUES (?, ?, ?, ?)",
                (username, password, role, full_name),
            )
            conn.commit()
            return True
        except sqlite3.IntegrityError:
            return False
        finally:
            conn.close()

    def update_user(self, user_id, username, password, role, full_name):
        conn = self.connect()
        try:
            conn.execute(
                "UPDATE users SET username = ?, password = ?, role = ?, full_name = ? WHERE id = ?",
                (username, password, role, full_name, user_id),
            )
            conn.commit()
            return True
        except sqlite3.IntegrityError:
            return False
        finally:
            conn.close()

    def delete_user(self, user_id):
        conn = self.connect()
        conn.execute("DELETE FROM users WHERE id = ?", (user_id,))
        conn.commit()
        conn.close()

    def add_defect(self, defect_data, created_by):
        conn = self.connect()
        conn.execute(
            """
            INSERT INTO defects (
                defect_code, model, line_name, defect_name, responsible, status,
                defect_date, description, root_cause, improvement_method,
                before_image_path, after_image_path, created_by,
                created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
            """,
            (
                defect_data.get("defect_code"),
                defect_data.get("model"),
                defect_data.get("line_name"),
                defect_data.get("defect_name"),
                defect_data.get("responsible"),
                defect_data.get("status"),
                defect_data.get("defect_date"),
                defect_data.get("description"),
                defect_data.get("root_cause"),
                defect_data.get("improvement_method"),
                defect_data.get("before_image_path"),
                defect_data.get("after_image_path"),
                created_by,
            ),
        )
        conn.commit()
        conn.close()

    def update_defect(self, defect_id, defect_data):
        conn = self.connect()
        conn.execute(
            """
            UPDATE defects SET
                defect_code = ?, model = ?, line_name = ?, defect_name = ?, responsible = ?,
                status = ?, defect_date = ?, description = ?, root_cause = ?,
                improvement_method = ?, before_image_path = ?, after_image_path = ?,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            (
                defect_data.get("defect_code"),
                defect_data.get("model"),
                defect_data.get("line_name"),
                defect_data.get("defect_name"),
                defect_data.get("responsible"),
                defect_data.get("status"),
                defect_data.get("defect_date"),
                defect_data.get("description"),
                defect_data.get("root_cause"),
                defect_data.get("improvement_method"),
                defect_data.get("before_image_path"),
                defect_data.get("after_image_path"),
                defect_id,
            ),
        )
        conn.commit()
        conn.close()

    def delete_defect(self, defect_id):
        conn = self.connect()
        conn.execute("DELETE FROM defects WHERE id = ?", (defect_id,))
        conn.commit()
        conn.close()

    def get_defects(self, search_text=""):
        conn = self.connect()
        if search_text:
            s = f"%{search_text}%"
            rows = conn.execute(
                """
                SELECT d.*, u.full_name as created_by_name
                FROM defects d
                LEFT JOIN users u ON u.id = d.created_by
                WHERE d.defect_name LIKE ? OR d.model LIKE ? OR d.responsible LIKE ? OR d.defect_code LIKE ?
                ORDER BY d.id DESC
                """,
                (s, s, s, s),
            ).fetchall()
        else:
            rows = conn.execute(
                """
                SELECT d.*, u.full_name as created_by_name
                FROM defects d
                LEFT JOIN users u ON u.id = d.created_by
                ORDER BY d.id DESC
                """
            ).fetchall()
        conn.close()
        return [dict(r) for r in rows]

    def get_defect_by_id(self, defect_id):
        conn = self.connect()
        row = conn.execute("SELECT * FROM defects WHERE id = ?", (defect_id,)).fetchone()
        conn.close()
        return dict(row) if row else None

    def get_monthly_summary(self):
        conn = self.connect()
        rows = conn.execute(
            """
            SELECT strftime('%Y-%m', defect_date) AS month_label, COUNT(*) AS total
            FROM defects
            WHERE defect_date IS NOT NULL
            GROUP BY strftime('%Y-%m', defect_date)
            ORDER BY month_label DESC
            """
        ).fetchall()
        conn.close()
        return [dict(r) for r in rows]

    def get_model_summary(self):
        conn = self.connect()
        rows = conn.execute(
            "SELECT model, COUNT(*) AS total FROM defects WHERE model IS NOT NULL AND model != '' GROUP BY model ORDER BY total DESC"
        ).fetchall()
        conn.close()
        return [dict(r) for r in rows]

    def get_responsible_summary(self):
        conn = self.connect()
        rows = conn.execute(
            "SELECT responsible, COUNT(*) AS total FROM defects WHERE responsible IS NOT NULL AND responsible != '' GROUP BY responsible ORDER BY total DESC"
        ).fetchall()
        conn.close()
        return [dict(r) for r in rows]

    def get_improvement_summary(self):
        conn = self.connect()
        rows = conn.execute(
            """
            SELECT improvement_method, COUNT(*) AS total
            FROM defects
            WHERE improvement_method IS NOT NULL AND improvement_method != ''
            GROUP BY improvement_method
            ORDER BY total DESC
            """
        ).fetchall()
        conn.close()
        return [dict(r) for r in rows]

    def export_csv(self, path):
        defects = self.get_defects()
        fieldnames = [
            "id", "defect_code", "model", "line_name", "defect_name", "responsible",
            "status", "defect_date", "description", "root_cause", "improvement_method",
            "before_image_path", "after_image_path", "created_at", "updated_at"
        ]
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            for item in defects:
                writer.writerow({k: item.get(k, "") for k in fieldnames})

    def export_excel(self, path):
        defects = self.get_defects()
        wb = Workbook()
        ws = wb.active
        ws.title = "Defects"

        headers = [
            "ID", "Mã lỗi", "Model", "Dây chuyền", "Tên lỗi", "Người phụ trách",
            "Trạng thái", "Ngày lỗi", "Mô tả", "Nguyên nhân", "Phương pháp cải thiện",
            "Ảnh trước", "Ảnh sau", "Ngày tạo", "Ngày cập nhật"
        ]
        ws.append(headers)

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row_num, item in enumerate(defects, start=2):
            ws[f"A{row_num}"] = item.get("id")
            ws[f"B{row_num}"] = item.get("defect_code")
            ws[f"C{row_num}"] = item.get("model")
            ws[f"D{row_num}"] = item.get("line_name")
            ws[f"E{row_num}"] = item.get("defect_name")
            ws[f"F{row_num}"] = item.get("responsible")
            ws[f"G{row_num}"] = item.get("status")
            ws[f"H{row_num}"] = item.get("defect_date")
            ws[f"I{row_num}"] = item.get("description")
            ws[f"J{row_num}"] = item.get("root_cause")
            ws[f"K{row_num}"] = item.get("improvement_method")
            ws[f"N{row_num}"] = item.get("created_at")
            ws[f"O{row_num}"] = item.get("updated_at")

            ws[f"I{row_num}"].alignment = Alignment(wrap_text=True, vertical="top")
            ws[f"J{row_num}"].alignment = Alignment(wrap_text=True, vertical="top")
            ws[f"K{row_num}"].alignment = Alignment(wrap_text=True, vertical="top")

            before = item.get("before_image_path")
            if before and os.path.exists(before):
                try:
                    img = XLImage(before)
                    img.width = 120
                    img.height = 90
                    ws.add_image(img, f"L{row_num}")
                except Exception:
                    ws[f"L{row_num}"] = "Ảnh lỗi"

            after = item.get("after_image_path")
            if after and os.path.exists(after):
                try:
                    img = XLImage(after)
                    img.width = 120
                    img.height = 90
                    ws.add_image(img, f"M{row_num}")
                except Exception:
                    ws[f"M{row_num}"] = "Ảnh lỗi"

            ws.row_dimensions[row_num].height = 90

        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 14
        ws.column_dimensions["C"].width = 16
        ws.column_dimensions["D"].width = 16
        ws.column_dimensions["E"].width = 22
        ws.column_dimensions["F"].width = 18
        ws.column_dimensions["G"].width = 14
        ws.column_dimensions["H"].width = 14
        ws.column_dimensions["I"].width = 36
        ws.column_dimensions["J"].width = 36
        ws.column_dimensions["K"].width = 28
        ws.column_dimensions["L"].width = 18
        ws.column_dimensions["M"].width = 18
        ws.column_dimensions["N"].width = 18
        ws.column_dimensions["O"].width = 18

        wb.save(path)

    def backup_database(self):
        if not os.path.exists(self.db_path):
            return ""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_file = BACKUP_DIR / f"error_tracker_backup_{timestamp}.db"
        shutil.copy2(self.db_path, backup_file)
        return str(backup_file)


def copy_image_to_storage(source_path, prefix):
    if not source_path or not os.path.exists(source_path):
        return ""
    file_name = f"{prefix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{os.path.basename(source_path)}"
    target_path = IMAGE_DIR / file_name
    shutil.copy2(source_path, target_path)
    return str(target_path)


def show_image_preview(label_widget, image_path):
    label_widget.configure(text="Không có ảnh", image=None)
    if not image_path or not os.path.exists(image_path):
        return
    try:
        img = PILImage.open(image_path).convert("RGB")
        img.thumbnail((220, 180))
        photo = ImageTk.PhotoImage(img)
        label_widget.configure(image=photo, text="")
        label_widget.image = photo
    except Exception:
        label_widget.configure(text=f"Ảnh: {os.path.basename(image_path)}")


def open_image_dialog():
    return filedialog.askopenfilename(
        title="Chọn ảnh",
        filetypes=[("Image files", "*.jpg *.jpeg *.png *.bmp *.gif")]
    )


class LoginWindow(ctk.CTk):
    def __init__(self, db):
        super().__init__()
        self.db = db
        self.title("Quality Error Tracker")
        self.geometry("440x320")
        self.minsize(440, 320)

        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("dark-blue")

        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(0, weight=1)

        container = ctk.CTkFrame(self, corner_radius=18)
        container.grid(row=0, column=0, padx=24, pady=24, sticky="nsew")

        title = ctk.CTkLabel(container, text="WELCOME TO ISSUE SYSTEM", font=ctk.CTkFont(size=24, weight="bold"))
        title.grid(row=0, column=0, columnspan=2, pady=(22, 16))

        self.username_var = tk.StringVar()
        self.password_var = tk.StringVar()

        ctk.CTkLabel(container, text="ACCOUNT").grid(row=1, column=0, sticky="w", padx=18, pady=(8, 4))
        self.entry_user = ctk.CTkEntry(container, textvariable=self.username_var, width=220)
        self.entry_user.grid(row=1, column=1, padx=(0, 18), pady=(8, 4))

        ctk.CTkLabel(container, text="PASSWORD").grid(row=2, column=0, sticky="w", padx=18, pady=(8, 4))
        self.entry_pass = ctk.CTkEntry(container, textvariable=self.password_var, width=220, show="*")
        self.entry_pass.grid(row=2, column=1, padx=(0, 18), pady=(8, 4))

        btn = ctk.CTkButton(container, text="LOGIN", command=self.login, width=220, height=40)
        btn.grid(row=3, column=0, columnspan=2, pady=(20, 8))

        self.bind("<Return>", lambda event: self.login())

    def login(self):
        username = self.username_var.get().strip()
        password = self.password_var.get()

        if not username:
            messagebox.showwarning("Thông báo", "Vui lòng nhập Username")
            return
        if not password:
            messagebox.showwarning("Thông báo", "Vui lòng nhập Password")
            return

        user = self.db.authenticate(username, password)
        if not user:
            messagebox.showerror("Lỗi", "Sai tài khoản hoặc mật khẩu")
            return

        self.destroy()
        app = MainApp(self.db, user)
        app.mainloop()


class UserFormDialog(ctk.CTkToplevel):
    def __init__(self, db, callback, user=None):
        super().__init__()
        self.db = db
        self.callback = callback
        self.user = user
        self.title("Thêm/Sửa tài khoản" if not user else "Cập nhật tài khoản")
        self.geometry("430x300")
        self.minsize(430, 300)

        self.grid_columnconfigure(0, weight=1)
        self.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(self, text="Username").grid(row=0, column=0, sticky="w", padx=18, pady=(14, 8))
        self.username_var = tk.StringVar(value=user["username"] if user else "")
        ctk.CTkEntry(self, textvariable=self.username_var, width=200).grid(row=0, column=1, padx=(0, 18), pady=(14, 8))

        ctk.CTkLabel(self, text="Password").grid(row=1, column=0, sticky="w", padx=18, pady=8)
        self.password_var = tk.StringVar(value=user["password"] if user else "")
        ctk.CTkEntry(self, textvariable=self.password_var, width=200, show="*").grid(row=1, column=1, padx=(0, 18), pady=8)

        ctk.CTkLabel(self, text="Role").grid(row=2, column=0, sticky="w", padx=18, pady=8)
        self.role_var = tk.StringVar(value=user["role"] if user else "Viewer")
        ctk.CTkComboBox(self, values=["Admin", "Engineer", "Viewer"], variable=self.role_var, width=200).grid(row=2, column=1, padx=(0, 18), pady=8)

        ctk.CTkLabel(self, text="Họ tên").grid(row=3, column=0, sticky="w", padx=18, pady=8)
        self.full_name_var = tk.StringVar(value=user["full_name"] if user else "")
        ctk.CTkEntry(self, textvariable=self.full_name_var, width=200).grid(row=3, column=1, padx=(0, 18), pady=8)

        ctk.CTkButton(self, text="Lưu", command=self.save, width=200).grid(row=4, column=0, columnspan=2, pady=(20, 8))

    def save(self):
        username = self.username_var.get().strip()
        password = self.password_var.get().strip()
        role = self.role_var.get()
        full_name = self.full_name_var.get().strip()

        if not username or not password or not role:
            messagebox.showerror("Lỗi", "Vui lòng điền đầy đủ thông tin")
            return

        if self.user:
            ok = self.db.update_user(self.user["id"], username, password, role, full_name)
        else:
            ok = self.db.insert_user(username, password, role, full_name)

        if not ok:
            messagebox.showerror("Lỗi", "Tên đăng nhập đã tồn tại hoặc dữ liệu không hợp lệ")
            return

        self.callback()
        self.destroy()


class MainApp(ctk.CTk):
    def __init__(self, db, user):
        super().__init__()
        self.db = db
        self.user = user
        self.role = user["role"]
        self.current_defect_id = None

        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("dark-blue")

        self.title("Quality Error Tracker")
        self.geometry("1400x840")
        self.minsize(1200, 760)

        self.grid_columnconfigure(0, weight=0)
        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=1)

        self.sidebar = ctk.CTkFrame(self, width=220, corner_radius=0)
        self.sidebar.grid(row=0, column=0, sticky="ns")

        self.content = ctk.CTkFrame(self, corner_radius=18)
        self.content.grid(row=0, column=1, padx=14, pady=14, sticky="nsew")

        self.build_sidebar()
        self.build_content()

        self.protocol("WM_DELETE_WINDOW", self.destroy)

    def build_sidebar(self):
        self.sidebar.grid_columnconfigure(0, weight=1)

        user_label = ctk.CTkLabel(
            self.sidebar,
            text=f"{self.user['full_name']}\n({self.user['role']})",
            font=ctk.CTkFont(size=18, weight="bold"),
            padx=12,
            pady=16
        )
        user_label.grid(row=0, column=0, pady=(16, 10), padx=12, sticky="ew")

        nav = [
            ("Danh sách lỗi", "list"),
            ("Thêm / Sửa lỗi", "form"),
            ("Dashboard", "dashboard"),
        ]
        if self.role == "Admin":
            nav.append(("Quản lý tài khoản", "account"))

        self.nav_buttons = {}

        for idx, (text, key) in enumerate(nav, start=1):
            btn = ctk.CTkButton(
                self.sidebar,
                text=text,
                command=lambda k=key: self.switch_tab(k),
                height=42,
                corner_radius=10
            )
            btn.grid(row=idx, column=0, padx=12, pady=6, sticky="ew")
            self.nav_buttons[key] = btn

        logout_btn = ctk.CTkButton(self.sidebar, text="Đăng xuất", command=self.logout, fg_color="#B00020", hover_color="#D32F2F", height=42)
        logout_btn.grid(row=len(nav) + 1, column=0, padx=12, pady=(18, 12), sticky="ew")

    def switch_tab(self, key):
        tab_map = {
            "list": self.list_frame,
            "form": self.form_frame,
            "dashboard": self.dashboard_frame,
            "account": self.account_frame
        }
        for k, frame in tab_map.items():
            frame.grid_forget()
        if key in tab_map:
            tab_map[key].grid(row=0, column=0, sticky="nsew", padx=10, pady=10)

    def build_content(self):
        self.content.grid_columnconfigure(0, weight=1)
        self.content.grid_rowconfigure(0, weight=1)

        self.list_frame = ctk.CTkFrame(self.content, corner_radius=14)
        self.form_frame = ctk.CTkFrame(self.content, corner_radius=14)
        self.dashboard_frame = ctk.CTkFrame(self.content, corner_radius=14)
        self.account_frame = ctk.CTkFrame(self.content, corner_radius=14)

        self.build_list_tab()
        self.build_form_tab()
        self.build_dashboard_tab()
        if self.role == "Admin":
            self.build_account_tab()

        self.switch_tab("list")

    def build_list_tab(self):
        self.list_frame.grid_columnconfigure(0, weight=1)
        self.list_frame.grid_rowconfigure(2, weight=1)

        top = ctk.CTkFrame(self.list_frame, corner_radius=12)
        top.grid(row=0, column=0, padx=12, pady=12, sticky="ew")

        self.search_var = ctk.CTkEntry(top, placeholder_text="Tìm kiếm theo mã lỗi, model, người phụ trách...")
        self.search_var.grid(row=0, column=0, padx=12, pady=12, sticky="ew")

        ctk.CTkButton(top, text="Tìm", command=self.load_defects_table, width=80).grid(row=0, column=1, padx=(0, 6), pady=12)
        ctk.CTkButton(top, text="Reset", command=self.reset_search, fg_color="#455A64", hover_color="#546E7A", width=80).grid(row=0, column=2, padx=(0, 12), pady=12)

        action_bar = ctk.CTkFrame(self.list_frame, corner_radius=12)
        action_bar.grid(row=1, column=0, padx=12, pady=(0, 12), sticky="ew")

        ctk.CTkButton(action_bar, text="Thêm mới", command=self.reset_form, width=120).grid(row=0, column=0, padx=10, pady=10)
        ctk.CTkButton(action_bar, text="Sửa", command=self.edit_selected_fault, width=120, state="normal" if self.role != "Viewer" else "disabled").grid(row=0, column=1, padx=10, pady=10)
        ctk.CTkButton(action_bar, text="Xóa", command=self.delete_selected_fault, width=120, fg_color="#D32F2F", hover_color="#B71C1C", state="normal" if self.role == "Admin" else "disabled").grid(row=0, column=2, padx=10, pady=10)
        ctk.CTkButton(action_bar, text="Export CSV", command=self.export_csv, width=120).grid(row=0, column=3, padx=10, pady=10)
        ctk.CTkButton(action_bar, text="Export Excel", command=self.export_excel, width=120).grid(row=0, column=4, padx=10, pady=10)
        ctk.CTkButton(action_bar, text="Backup", command=self.backup_database, width=120).grid(row=0, column=5, padx=10, pady=10)

        table_container = ctk.CTkFrame(self.list_frame, corner_radius=12)
        table_container.grid(row=2, column=0, padx=12, pady=(0, 12), sticky="nsew")

        table_container.grid_columnconfigure(0, weight=1)
        table_container.grid_rowconfigure(0, weight=1)

        self.list_tree = tk.ttk.Treeview(table_container, columns=("id","defect_code","model","line_name","defect_name","responsible","status","defect_date"), show="headings", height=17)
        self.list_tree.grid(row=0, column=0, sticky="nsew", padx=10, pady=10)

        for col, text in zip(
            ("id","defect_code","model","line_name","defect_name","responsible","status","defect_date"),
            ("ID","Mã Lỗi","Model","Line","Tên Lỗi","Người Phụ Trách","Trạng Thái","Ngày")
        ):
            self.list_tree.heading(col, text=text)
            self.list_tree.column(col, width=110, anchor="center")

        self.list_tree.tag_configure("Open", background="#2E7DFF", foreground="white")
        self.list_tree.tag_configure("Processed", background="#7C4DFF", foreground="white")
        self.list_tree.tag_configure("Processing", background="#FFB300", foreground="black")
        self.list_tree.tag_configure("Closed", background="#2E7D32", foreground="white")

        self.list_tree.bind("<<TreeviewSelect>>", self.on_list_select)

    def build_form_tab(self):
        self.form_frame.grid_columnconfigure(0, weight=1)
        self.form_frame.grid_rowconfigure(0, weight=1)

        # Thêm scrollbar cho form
        canvas = tk.Canvas(self.form_frame, bg="#212121", highlightthickness=0)
        scrollbar = ttk.Scrollbar(self.form_frame, orient="vertical", command=canvas.yview)
        scrollable_frame = ctk.CTkFrame(canvas, fg_color="#212121")

        scrollable_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )

        canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)

        canvas.grid(row=0, column=0, sticky="nsew")
        scrollbar.grid(row=0, column=1, sticky="ns")

        self.form_frame.grid_columnconfigure(0, weight=1)
        self.form_frame.grid_rowconfigure(0, weight=1)

        form_wrap = scrollable_frame

        row = 0
        ctk.CTkLabel(form_wrap, text="Mã lỗi").grid(row=row, column=0, padx=12, pady=(18, 8), sticky="w")
        self.defect_code = ctk.CTkEntry(form_wrap, width=220)
        self.defect_code.grid(row=row, column=1, padx=12, pady=(18, 8))

        ctk.CTkLabel(form_wrap, text="Model").grid(row=row, column=2, padx=12, pady=(18, 8), sticky="w")
        self.model = ctk.CTkEntry(form_wrap, width=220)
        self.model.grid(row=row, column=3, padx=12, pady=(18, 8))

        row += 1
        ctk.CTkLabel(form_wrap, text="Line").grid(row=row, column=0, padx=12, pady=8, sticky="w")
        self.line_name = ctk.CTkEntry(form_wrap, width=220)
        self.line_name.grid(row=row, column=1, padx=12, pady=8)

        ctk.CTkLabel(form_wrap, text="Tên lỗi").grid(row=row, column=2, padx=12, pady=8, sticky="w")
        self.defect_name = ctk.CTkEntry(form_wrap, width=220)
        self.defect_name.grid(row=row, column=3, padx=12, pady=8)

        row += 1
        ctk.CTkLabel(form_wrap, text="Người phụ trách").grid(row=row, column=0, padx=12, pady=8, sticky="w")
        self.responsible = ctk.CTkEntry(form_wrap, width=220)
        self.responsible.grid(row=row, column=1, padx=12, pady=8)

        ctk.CTkLabel(form_wrap, text="Trạng thái").grid(row=row, column=2, padx=12, pady=8, sticky="w")
        self.status = ctk.CTkComboBox(form_wrap, values=["Open", "Processed", "Processing", "Closed"], width=200)
        self.status.grid(row=row, column=3, padx=12, pady=8)
        self.status.set("Open")

        row += 1
        ctk.CTkLabel(form_wrap, text="Ngày lỗi").grid(row=row, column=0, padx=12, pady=8, sticky="w")
        self.defect_date = ctk.CTkEntry(form_wrap, width=220)
        self.defect_date.grid(row=row, column=1, padx=12, pady=8)
        self.defect_date.insert("end", datetime.today().strftime("%Y-%m-%d"))

        ctk.CTkLabel(form_wrap, text="Phương pháp cải thiện").grid(row=row, column=2, padx=12, pady=8, sticky="w")
        self.improvement_method = ctk.CTkEntry(form_wrap, width=220)
        self.improvement_method.grid(row=row, column=3, padx=12, pady=8)

        row += 1
        ctk.CTkLabel(form_wrap, text="Mô tả").grid(row=row, column=0, padx=12, pady=(10, 8), sticky="nw")
        self.description = ctk.CTkTextbox(form_wrap, width=220, height=110)
        self.description.grid(row=row, column=1, padx=12, pady=(10, 8))
        self.description.bind("<Shift-Return>", lambda e: self.insert_newline(self.description))

        ctk.CTkLabel(form_wrap, text="Nguyên nhân").grid(row=row, column=2, padx=12, pady=(10, 8), sticky="nw")
        self.root_cause = ctk.CTkTextbox(form_wrap, width=220, height=110)
        self.root_cause.grid(row=row, column=3, padx=12, pady=(10, 8))
        self.root_cause.bind("<Shift-Return>", lambda e: self.insert_newline(self.root_cause))

        row += 1
        image_panel = ctk.CTkFrame(form_wrap, corner_radius=12)
        image_panel.grid(row=row, column=0, columnspan=4, padx=12, pady=(12, 18), sticky="nsew")

        image_panel.grid_columnconfigure((0, 1), weight=1)

        before_box = ctk.CTkFrame(image_panel, corner_radius=12)
        before_box.grid(row=0, column=0, padx=12, pady=12, sticky="nsew")
        ctk.CTkLabel(before_box, text="Ảnh trước").pack(padx=10, pady=(10, 4))
        self.before_label = ctk.CTkLabel(before_box, text="Không có ảnh", width=220, height=140)
        self.before_label.pack(padx=10, pady=(0, 8))
        ctk.CTkButton(before_box, text="Chọn ảnh trước", command=self.select_before_image).pack(padx=10, pady=(0, 10))

        after_box = ctk.CTkFrame(image_panel, corner_radius=12)
        after_box.grid(row=0, column=1, padx=12, pady=12, sticky="nsew")
        ctk.CTkLabel(after_box, text="Ảnh sau").pack(padx=10, pady=(10, 4))
        self.after_label = ctk.CTkLabel(after_box, text="Không có ảnh", width=220, height=140)
        self.after_label.pack(padx=10, pady=(0, 8))
        ctk.CTkButton(after_box, text="Chọn ảnh sau", command=self.select_after_image).pack(padx=10, pady=(0, 10))

        self.before_image_path = ""
        self.after_image_path = ""

        row += 1
        btn_row = ctk.CTkFrame(form_wrap, corner_radius=12)
        btn_row.grid(row=row, column=0, columnspan=4, padx=12, pady=(0, 18), sticky="ew")

        ctk.CTkButton(btn_row, text="Lưu dữ liệu", command=self.save_defect, width=170).pack(side="left", padx=12, pady=10)
        ctk.CTkButton(btn_row, text="Xóa form", command=self.reset_form, fg_color="#455A64", hover_color="#546E7A", width=170).pack(side="left", padx=12, pady=10)

    def insert_newline(self, widget):
        widget.insert("insert", "\n")
        return "break"

    def build_dashboard_tab(self):
        self.dashboard_frame.grid_columnconfigure(0, weight=1)
        self.dashboard_frame.grid_rowconfigure(1, weight=1)

        summary = ctk.CTkFrame(self.dashboard_frame, corner_radius=12)
        summary.grid(row=0, column=0, padx=12, pady=12, sticky="ew")

        summary.grid_columnconfigure((0, 1, 2), weight=1)

        self.total_label = ctk.CTkLabel(summary, text="Tổng lỗi: 0", font=ctk.CTkFont(size=18, weight="bold"))
        self.total_label.grid(row=0, column=0, padx=12, pady=12)

        self.month_label = ctk.CTkLabel(summary, text="Tháng nhiều nhất: -", font=ctk.CTkFont(size=14))
        self.month_label.grid(row=0, column=1, padx=12, pady=12)

        self.model_label = ctk.CTkLabel(summary, text="Model phổ biến: -", font=ctk.CTkFont(size=14))
        self.model_label.grid(row=0, column=2, padx=12, pady=12)

        chart_box = ctk.CTkFrame(self.dashboard_frame, corner_radius=12)
        chart_box.grid(row=1, column=0, padx=12, pady=(0, 12), sticky="nsew")

        self.chart_box = chart_box
        self.chart_canvas = None

    def build_account_tab(self):
        self.account_frame.grid_columnconfigure(0, weight=1)
        self.account_frame.grid_rowconfigure(1, weight=1)

        toolbar = ctk.CTkFrame(self.account_frame, corner_radius=12)
        toolbar.grid(row=0, column=0, padx=12, pady=12, sticky="ew")

        ctk.CTkButton(toolbar, text="Thêm tài khoản", command=self.open_user_form, width=160).pack(side="left", padx=10, pady=10)
        ctk.CTkButton(toolbar, text="Sửa tài khoản", command=self.edit_selected_user, width=160).pack(side="left", padx=10, pady=10)
        ctk.CTkButton(toolbar, text="Xóa tài khoản", command=self.delete_selected_user, fg_color="#D32F2F", hover_color="#B71C1C", width=160).pack(side="left", padx=10, pady=10)

        table = ctk.CTkFrame(self.account_frame, corner_radius=12)
        table.grid(row=1, column=0, padx=12, pady=(0, 12), sticky="nsew")

        table.grid_columnconfigure(0, weight=1)
        table.grid_rowconfigure(0, weight=1)

        self.user_tree = tk.ttk.Treeview(table, columns=("id","username","role","full_name"), show="headings", height=18)
        self.user_tree.grid(row=0, column=0, sticky="nsew", padx=10, pady=10)

        self.user_tree.heading("id", text="ID")
        self.user_tree.heading("username", text="Username")
        self.user_tree.heading("role", text="Role")
        self.user_tree.heading("full_name", text="Họ tên")

        for col in ("id","username","role","full_name"):
            self.user_tree.column(col, width=150, anchor="center")

        self.load_users_table()

    def reset_search(self):
        self.search_var.delete(0, "end")
        self.load_defects_table()

    def load_defects_table(self):
        for item in self.list_tree.get_children():
            self.list_tree.delete(item)

        rows = self.db.get_defects(self.search_var.get().strip())
        for item in rows:
            status = item.get("status", "Open")
            self.list_tree.insert(
                "",
                "end",
                iid=str(item["id"]),
                values=(
                    item.get("id"), item.get("defect_code"), item.get("model"), item.get("line_name"),
                    item.get("defect_name"), item.get("responsible"), status, item.get("defect_date")
                ),
                tags=(status,),
            )

        for row_id in self.list_tree.get_children():
            status = self.list_tree.item(row_id, "values")[6]
            self.list_tree.item(row_id, tags=(status,))

    def on_list_select(self, event):
        selected = self.list_tree.selection()
        if not selected:
            return
        item = self.list_tree.item(selected[0], "values")
        defect_id = item[0]
        self.current_defect_id = defect_id
        defect = self.db.get_defect_by_id(defect_id)
        if defect:
            self.fill_form(defect)

    def fill_form(self, defect):
        self.defect_code.delete(0, "end")
        self.model.delete(0, "end")
        self.line_name.delete(0, "end")
        self.defect_name.delete(0, "end")
        self.responsible.delete(0, "end")
        self.defect_date.delete(0, "end")
        self.improvement_method.delete(0, "end")

        self.defect_code.insert("end", defect.get("defect_code", ""))
        self.model.insert("end", defect.get("model", ""))
        self.line_name.insert("end", defect.get("line_name", ""))
        self.defect_name.insert("end", defect.get("defect_name", ""))
        self.responsible.insert("end", defect.get("responsible", ""))
        self.defect_date.insert("end", defect.get("defect_date", ""))
        self.improvement_method.insert("end", defect.get("improvement_method", ""))

        self.status.set(defect.get("status", "Open"))

        self.description.delete("0.0", "end")
        self.root_cause.delete("0.0", "end")
        self.description.insert("0.0", defect.get("description", ""))
        self.root_cause.insert("0.0", defect.get("root_cause", ""))

        self.before_image_path = defect.get("before_image_path", "") or ""
        self.after_image_path = defect.get("after_image_path", "") or ""
        show_image_preview(self.before_label, self.before_image_path)
        show_image_preview(self.after_label, self.after_image_path)

    def reset_form(self):
        self.current_defect_id = None
        self.defect_code.delete(0, "end")
        self.model.delete(0, "end")
        self.line_name.delete(0, "end")
        self.defect_name.delete(0, "end")
        self.responsible.delete(0, "end")
        self.status.set("Open")
        self.defect_date.delete(0, "end")
        self.improvement_method.delete(0, "end")

        self.defect_date.insert("end", datetime.today().strftime("%Y-%m-%d"))

        self.description.delete("0.0", "end")
        self.root_cause.delete("0.0", "end")

        # ✅ FIX: Xóa hình ảnh cũ
        self.before_image_path = ""
        self.after_image_path = ""
        show_image_preview(self.before_label, "")
        show_image_preview(self.after_label, "")

    def select_before_image(self):
        selected = open_image_dialog()
        if selected:
            self.before_image_path = selected
            show_image_preview(self.before_label, selected)

    def select_after_image(self):
        selected = open_image_dialog()
        if selected:
            self.after_image_path = selected
            show_image_preview(self.after_label, selected)

    def save_defect(self):
        if self.role == "Viewer":
            messagebox.showerror("Lỗi", "Viewer không có quyền lưu dữ liệu")
            return

        defect_data = {
            "defect_code": self.defect_code.get().strip(),
            "model": self.model.get().strip(),
            "line_name": self.line_name.get().strip(),
            "defect_name": self.defect_name.get().strip(),
            "responsible": self.responsible.get().strip(),
            "status": self.status.get(),
            "defect_date": self.defect_date.get().strip(),
            "description": self.description.get("0.0", "end").strip(),
            "root_cause": self.root_cause.get("0.0", "end").strip(),
            "improvement_method": self.improvement_method.get().strip(),
            "before_image_path": self.before_image_path,
            "after_image_path": self.after_image_path,
        }

        if not defect_data["defect_name"] or not defect_data["model"] or not defect_data["responsible"]:
            messagebox.showerror("Lỗi", "Vui lòng điền tên lỗi, model và người phụ trách")
            return

        if self.before_image_path:
            defect_data["before_image_path"] = copy_image_to_storage(self.before_image_path, "before")
        if self.after_image_path:
            defect_data["after_image_path"] = copy_image_to_storage(self.after_image_path, "after")

        if self.current_defect_id:
            self.db.update_defect(self.current_defect_id, defect_data)
            messagebox.showinfo("Thông báo", "Cập nhật thành công")
        else:
            self.db.add_defect(defect_data, self.user["id"])
            messagebox.showinfo("Thông báo", "Thêm dữ liệu thành công")

        self.load_defects_table()
        self.load_dashboard()
        self.reset_form()

    def edit_selected_fault(self):
        if self.role == "Viewer":
            messagebox.showerror("Lỗi", "Viewer không có quyền sửa dữ liệu")
            return
        if not self.current_defect_id:
            messagebox.showwarning("Cảnh báo", "Vui lòng chọn một lỗi trong danh sách")
            return
        defect = self.db.get_defect_by_id(self.current_defect_id)
        if defect:
            self.fill_form(defect)
            self.switch_tab("form")

    def delete_selected_fault(self):
        if self.role != "Admin":
            messagebox.showerror("Lỗi", "Bạn không có quyền xóa dữ liệu")
            return
        if not self.current_defect_id:
            messagebox.showwarning("Cảnh báo", "Vui lòng chọn một lỗi để xóa")
            return
        if messagebox.askyesno("Xác nhận", "Bạn có chắc chắn muốn xóa dữ liệu này?"):
            self.db.delete_defect(self.current_defect_id)
            self.current_defect_id = None
            self.reset_form()
            self.load_defects_table()
            self.load_dashboard()

    def export_csv(self):
        if self.role == "Viewer":
            messagebox.showerror("Lỗi", "Viewer không có quyền xuất dữ liệu")
            return
        target = EXPORT_DIR / f"defects_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        try:
            self.db.export_csv(str(target))
            messagebox.showinfo("Thành công", f"Đã xuất CSV: {target}")
        except Exception as e:
            messagebox.showerror("Lỗi", f"Không thể xuất CSV: {e}")

    def export_excel(self):
        if self.role == "Viewer":
            messagebox.showerror("Lỗi", "Viewer không có quyền xuất báo cáo Excel")
            return
        target = EXPORT_DIR / f"defects_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        try:
            self.db.export_excel(str(target))
            messagebox.showinfo("Thành công", f"Đã xuất Excel: {target}")
        except Exception as e:
            messagebox.showerror("Lỗi", f"Không thể xuất Excel: {e}")

    def backup_database(self):
        if self.role == "Viewer":
            messagebox.showerror("Lỗi", "Viewer không có quyền backup dữ liệu")
            return
        path = self.db.backup_database()
        if path:
            messagebox.showinfo("Thành công", f"Backup thành công: {path}")
        else:
            messagebox.showerror("Lỗi", "Không tìm thấy database để backup")

    def draw_method_chart(self):
        data = self.db.get_improvement_summary()
        if not data:
            return

        methods = [x["improvement_method"] for x in data]
        totals = [x["total"] for x in data]
        total_count = sum(totals)
        percents = [round((t / total_count) * 100, 1) for t in totals]

        if self.chart_canvas is not None:
            self.chart_canvas.get_tk_widget().destroy()

        fig = Figure(figsize=(9, 4), dpi=100)
        ax = fig.add_subplot(111)

        bars = ax.bar(methods, percents, color="#4EB5FF")
        ax.set_title("Tỉ lệ phương pháp cải thiện (%)", fontsize=12, fontweight="bold")
        ax.set_xlabel("Phương pháp")
        ax.set_ylabel("Tỉ lệ %")

        for bar, pct in zip(bars, percents):
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                pct + 0.6,
                f"{pct}%",
                ha="center",
                va="bottom",
                fontsize=9,
                fontweight="bold",
            )

        fig.tight_layout()
        self.chart_canvas = FigureCanvasTkAgg(fig, master=self.chart_box)
        self.chart_canvas.draw()
        self.chart_canvas.get_tk_widget().pack(fill="both", expand=True)

    def load_dashboard(self):
        total_rows = self.db.get_defects()
        self.total_label.configure(text=f"Tổng lỗi: {len(total_rows)}")

        months = self.db.get_monthly_summary()
        if months:
            top_month = max(months, key=lambda x: x["total"])
            self.month_label.configure(text=f"Tháng nhiều nhất: {top_month.get('month_label', '-')}")
        else:
            self.month_label.configure(text="Tháng nhiều nhất: -")

        models = self.db.get_model_summary()
        if models:
            top_model = max(models, key=lambda x: x["total"])
            self.model_label.configure(text=f"Model phổ biến: {top_model.get('model', '-')}")
        else:
            self.model_label.configure(text="Model phổ biến: -")

        self.draw_method_chart()

    def open_user_form(self):
        UserFormDialog(self.db, self.load_users_table)

    def load_users_table(self):
        if not hasattr(self, "user_tree"):
            return
        for row in self.user_tree.get_children():
            self.user_tree.delete(row)

        for user in self.db.get_all_users():
            self.user_tree.insert("", "end", values=(user["id"], user["username"], user["role"], user["full_name"]))

    def edit_selected_user(self):
        sel = self.user_tree.selection()
        if not sel:
            messagebox.showwarning("Cảnh báo", "Vui lòng chọn tài khoản để sửa")
            return
        item = self.user_tree.item(sel[0], "values")
        user = self.db.get_user_by_id(item[0])
        if user:
            UserFormDialog(self.db, self.load_users_table, user=user)

    def delete_selected_user(self):
        sel = self.user_tree.selection()
        if not sel:
            messagebox.showwarning("Cảnh báo", "Vui lòng chọn tài khoản để xóa")
            return
        item = self.user_tree.item(sel[0], "values")
        user_id = item[0]
        if user_id == self.user["id"]:
            messagebox.showerror("Lỗi", "Không thể xóa tài khoản đang đăng nhập")
            return
        if messagebox.askyesno("Xác nhận", "Bạn có chắc chắn muốn xóa tài khoản này?"):
            self.db.delete_user(user_id)
            self.load_users_table()

    def logout(self):
        self.destroy()
        app = LoginWindow(Database())
        app.mainloop()


if __name__ == "__main__":
    db = Database()
    app = LoginWindow(db)
    app.mainloop()
